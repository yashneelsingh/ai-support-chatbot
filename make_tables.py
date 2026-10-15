"""Format List of Figures and List of Tables as actual tables in Internship_Report_v5.docx"""

import docx
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

REPORT = r'd:\ai-support-chatbot\Internship_Report_v5.docx'

doc = docx.Document(REPORT)

# We need to find the sections for LIST OF FIGURES and LIST OF TABLES
lof_start = None
lot_start = None
ch1_start = None

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    style = p.style.name if p.style else ''
    if text == 'LIST OF FIGURES' and 'Heading' in style:
        lof_start = i
    elif text == 'LIST OF TABLES' and 'Heading' in style:
        lot_start = i
    elif text == 'Chapter 1: Introduction' and 'Heading 1' in style:
        ch1_start = i
        break

if lof_start is None or lot_start is None or ch1_start is None:
    print("ERROR: Could not find required sections.")
    sys.exit(1)

print(f"Found sections: LOF={lof_start}, LOT={lot_start}, Ch1={ch1_start}")

# Extract figures
figures = []
for i in range(lof_start + 1, lot_start):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if text.startswith('Figure '):
        figures.append((p, text))

# Extract tables
tables = []
for i in range(lot_start + 1, ch1_start):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if text.startswith('Table '):
        tables.append((p, text))

print(f"Found {len(figures)} figures and {len(tables)} tables to format.")

# Function to extract label and description
def parse_caption(text):
    if ':' in text:
        parts = text.split(':', 1)
        return parts[0].strip(), parts[1].strip()
    return text, ""

# For LOF, insert a table right after the LOF heading
lof_heading = doc.paragraphs[lof_start]
# We will insert a table by adding it to the parent element, but python-docx doesn't easily allow inserting tables at arbitrary positions.
# The standard way is to add a table to the end and then move its XML element.

def insert_table_after_paragraph(paragraph, items, col1_width=1.5, col2_width=5.0):
    # Create a new table at the end of the document
    table = doc.add_table(rows=len(items), cols=2)
    table.style = 'Table Grid'
    table.autofit = False
    
    # Set column widths
    table.columns[0].width = Inches(col1_width)
    table.columns[1].width = Inches(col2_width)
    
    for row_idx, (p, text) in enumerate(items):
        label, desc = parse_caption(text)
        row = table.rows[row_idx]
        row.cells[0].text = label
        row.cells[1].text = desc
        
        # Set cell widths explicitly
        row.cells[0].width = Inches(col1_width)
        row.cells[1].width = Inches(col2_width)

    # Move the table XML right after the paragraph
    p_element = paragraph._element
    tbl_element = table._element
    p_element.addnext(tbl_element)
    
    # Remove the original paragraphs
    for p, _ in items:
        p_element = p._element
        parent = p_element.getparent()
        if parent is not None:
            parent.remove(p_element)

    return table

# Insert LOF table
insert_table_after_paragraph(lof_heading, figures)

# Insert LOT table
# We need to get the lot_heading again because indices might have changed, 
# but we stored the exact paragraph objects in `tables` which we can use or just find by text.
lot_heading = None
for p in doc.paragraphs:
    text = p.text.strip()
    style = p.style.name if p.style else ''
    if text == 'LIST OF TABLES' and 'Heading' in style:
        lot_heading = p
        break

if lot_heading:
    insert_table_after_paragraph(lot_heading, tables)

doc.save(REPORT)
print("Successfully converted List of Figures and List of Tables to tables.")
